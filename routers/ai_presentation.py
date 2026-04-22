"""AI presentation routes (teacher + admin)."""
import asyncio
import json
import logging
import re
import traceback
from datetime import date
from io import BytesIO
from typing import Optional
from urllib.parse import quote

import httpx

log = logging.getLogger("ai_presentation")
from fastapi import APIRouter, Request, Depends, Form, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse, JSONResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy import select, delete, func
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from models import get_db, Employee, Subject, AIPresentation, School, AIQuestionSet
from dependencies import require_login, require_admin, flash, get_template_context, page_info
from ai_service import generate_presentation
from language import language_manager


def _t(request: Request, key: str) -> str:
    lang = request.session.get('language', 'uz')
    return language_manager.get(key, lang)

router = APIRouter()
templates = Jinja2Templates(directory="templates")

LANGUAGE_DISPLAY = {
    'uz': "O'zbekcha",
    'ru': "Русский",
    'en': "English",
}

TEMPLATES: dict = {
    'cosmos': {
        'label': 'Cosmos',
        'stage_bg': 'linear-gradient(135deg, #182448, #0a0e1c)',
        'accent': '#64c8ff',
        'pptx': {
            'title':    (0x1a, 0x4d, 0x9e),
            'subtitle': (0x55, 0x66, 0x77),
            'muted':    (0x88, 0x88, 0x88),
            'correct':  (0x1d, 0x9b, 0x4f),
            'wrong':    (0xc0, 0x39, 0x2b),
            'body':     (0x33, 0x33, 0x33),
        },
    },
    'ocean': {
        'label': 'Ocean',
        'stage_bg': 'linear-gradient(135deg, #0d2b3e, #071a2c)',
        'accent': '#00d4aa',
        'pptx': {
            'title':    (0x0d, 0x5c, 0x6e),
            'subtitle': (0x3d, 0x7a, 0x8a),
            'muted':    (0x70, 0x90, 0x95),
            'correct':  (0x00, 0xb8, 0x8a),
            'wrong':    (0xc0, 0x39, 0x2b),
            'body':     (0x1a, 0x3a, 0x40),
        },
    },
    'aurora': {
        'label': 'Aurora',
        'stage_bg': 'linear-gradient(135deg, #1e1244, #0d0a2e)',
        'accent': '#b478ff',
        'pptx': {
            'title':    (0x5a, 0x1e, 0x9e),
            'subtitle': (0x7a, 0x5a, 0xaa),
            'muted':    (0x88, 0x70, 0xa0),
            'correct':  (0x1d, 0x9b, 0x4f),
            'wrong':    (0xc0, 0x39, 0x2b),
            'body':     (0x2a, 0x1a, 0x44),
        },
    },
}


# ============================================================================
# TEACHER ROUTES
# ============================================================================

@router.get("/teacher/ai", response_class=HTMLResponse, dependencies=[Depends(require_login)])
async def teacher_ai_list(request: Request, page: int = 1, db: AsyncSession = Depends(get_db)):
    per_page = 6
    teacher_id = request.session.get('user_id')
    teacher = (await db.execute(select(Employee).where(Employee.id == teacher_id))).scalar_one()

    if not teacher.ai_enabled:
        flash(request, _t(request, 'ai_disabled_for_you'), 'warning')
        return RedirectResponse(url="/teacher/dashboard", status_code=303)

    _reset_if_new_day(teacher)
    await db.commit()

    total = (await db.execute(
        select(func.count(AIPresentation.id)).where(AIPresentation.teacher_id == teacher_id)
    )).scalar()
    pg = page_info(total, page, per_page, request)

    presentations = (await db.execute(
        select(AIPresentation)
        .where(AIPresentation.teacher_id == teacher_id)
        .order_by(AIPresentation.created_at.desc())
        .offset(pg['row_offset']).limit(per_page)
    )).scalars().all()

    context = await get_template_context(request, db)
    context.update({
        'presentations': presentations,
        'teacher': teacher,
        'remaining': max(0, teacher.ai_daily_limit - teacher.ai_used_today),
        'language_display': LANGUAGE_DISPLAY,
        'presentation_templates': TEMPLATES,
        **pg,
    })
    return templates.TemplateResponse('teacher/ai_list.html', context)


@router.get("/teacher/ai/create", response_class=HTMLResponse, dependencies=[Depends(require_login)])
async def teacher_ai_create_form(request: Request, db: AsyncSession = Depends(get_db)):
    teacher_id = request.session.get('user_id')
    teacher = (await db.execute(select(Employee).where(Employee.id == teacher_id))).scalar_one()

    if not teacher.ai_enabled:
        flash(request, _t(request, 'ai_disabled_short'), 'warning')
        return RedirectResponse(url="/teacher/dashboard", status_code=303)

    _reset_if_new_day(teacher)
    await db.commit()

    if teacher.ai_used_today >= teacher.ai_daily_limit:
        flash(request, _t(request, 'ai_limit_reached_msg'), 'warning')
        return RedirectResponse(url="/teacher/ai", status_code=303)

    subjects = (await db.execute(select(Subject).order_by(Subject.name))).scalars().all()

    context = await get_template_context(request, db)
    context.update({
        'subjects': subjects,
        'remaining': teacher.ai_daily_limit - teacher.ai_used_today,
        'templates': TEMPLATES,
    })
    return templates.TemplateResponse('teacher/ai_create.html', context)


@router.post("/teacher/ai/generate", dependencies=[Depends(require_login)])
async def teacher_ai_generate(
    request: Request,
    subject_id: int = Form(...),
    grade: int = Form(...),
    topic: str = Form(...),
    language: str = Form('uz'),
    template: str = Form('cosmos'),
    db: AsyncSession = Depends(get_db),
):
    teacher_id = request.session.get('user_id')
    teacher = (await db.execute(select(Employee).where(Employee.id == teacher_id))).scalar_one()

    if not teacher.ai_enabled:
        return JSONResponse({'ok': False, 'error': _t(request, 'ai_disabled_short')}, status_code=403)

    _reset_if_new_day(teacher)

    if teacher.ai_used_today >= teacher.ai_daily_limit:
        return JSONResponse({'ok': False, 'error': _t(request, 'ai_limit_reached_short')}, status_code=429)

    if not (1 <= grade <= 11):
        return JSONResponse({'ok': False, 'error': _t(request, 'ai_invalid_grade')}, status_code=400)

    topic = (topic or '').strip()
    if not topic or len(topic) > 300:
        return JSONResponse({'ok': False, 'error': _t(request, 'ai_invalid_topic')}, status_code=400)

    if language not in ('uz', 'ru', 'en'):
        language = 'uz'

    subject = (await db.execute(select(Subject).where(Subject.id == subject_id))).scalar_one_or_none()
    if not subject:
        return JSONResponse({'ok': False, 'error': _t(request, 'ai_subject_not_found')}, status_code=400)

    try:
        data = await generate_presentation(subject.name, grade, topic, language)
    except Exception as e:
        return JSONResponse({'ok': False, 'error': f"{_t(request, 'ai_error_prefix')}: {str(e)[:200]}"}, status_code=500)

    tpl_key = template if template in TEMPLATES else 'cosmos'

    presentation = AIPresentation(
        teacher_id=teacher_id,
        subject_id=subject.id,
        subject_name=subject.name,
        grade=grade,
        topic=topic,
        language=language,
        template=tpl_key,
        content=json.dumps(data, ensure_ascii=False),
        slides_count=len(data.get('slides', [])),
    )
    db.add(presentation)

    teacher.ai_used_today += 1
    teacher.ai_last_reset = date.today()

    await db.commit()
    await db.refresh(presentation)

    return JSONResponse({'ok': True, 'id': presentation.id, 'redirect': f"/teacher/ai/{presentation.id}"})


@router.get("/teacher/ai/{pid}", response_class=HTMLResponse, dependencies=[Depends(require_login)])
async def teacher_ai_view(pid: int, request: Request, db: AsyncSession = Depends(get_db)):
    teacher_id = request.session.get('user_id')
    presentation = (await db.execute(
        select(AIPresentation).where(AIPresentation.id == pid)
    )).scalar_one_or_none()

    if not presentation:
        raise HTTPException(404)

    is_admin = request.session.get('is_admin', False)
    if not is_admin and presentation.teacher_id != teacher_id:
        raise HTTPException(403)

    data = json.loads(presentation.content)

    tpl_key = presentation.template if presentation.template in TEMPLATES else 'cosmos'

    context = await get_template_context(request, db)
    context.update({
        'presentation': presentation,
        'data': data,
        'language_display': LANGUAGE_DISPLAY,
        'allow_download': is_admin,
        'template_info': TEMPLATES[tpl_key],
        'template_key': tpl_key,
    })
    return templates.TemplateResponse('teacher/ai_view.html', context)


@router.post("/teacher/ai/{pid}/delete", dependencies=[Depends(require_login)])
async def teacher_ai_delete(pid: int, request: Request, db: AsyncSession = Depends(get_db)):
    teacher_id = request.session.get('user_id')
    presentation = (await db.execute(
        select(AIPresentation).where(AIPresentation.id == pid)
    )).scalar_one_or_none()

    if not presentation:
        raise HTTPException(404)

    is_admin = request.session.get('is_admin', False)
    if not is_admin and presentation.teacher_id != teacher_id:
        raise HTTPException(403)

    await db.delete(presentation)
    await db.commit()
    flash(request, _t(request, 'ai_presentation_deleted'), 'success')

    redirect = "/admin/ai" if is_admin else "/teacher/ai"
    return RedirectResponse(url=redirect, status_code=303)


# ============================================================================
# ADMIN ROUTES
# ============================================================================

@router.get("/admin/ai", response_class=HTMLResponse, dependencies=[Depends(require_admin)])
async def admin_ai_list(
    request: Request,
    teacher_id: Optional[str] = None,
    page: int = 1,
    db: AsyncSession = Depends(get_db),
):
    per_page = 6
    tid: Optional[int] = int(teacher_id) if teacher_id and teacher_id.strip().isdigit() else None

    count_q = select(func.count(AIPresentation.id))
    data_q = (
        select(AIPresentation, Employee)
        .join(Employee, Employee.id == AIPresentation.teacher_id)
        .order_by(AIPresentation.created_at.desc())
    )
    if tid:
        count_q = count_q.where(AIPresentation.teacher_id == tid)
        data_q = data_q.where(AIPresentation.teacher_id == tid)

    total = (await db.execute(count_q)).scalar()
    pg = page_info(total, page, per_page, request)

    rows = (await db.execute(data_q.offset(pg['row_offset']).limit(per_page))).all()
    items = [{'p': p, 't': t} for p, t in rows]

    teachers = (await db.execute(
        select(Employee).where(Employee.is_admin.is_(False)).order_by(Employee.first_name)
    )).scalars().all()

    context = await get_template_context(request, db)
    context.update({
        'items': items,
        'teachers': teachers,
        'selected_teacher_id': tid,
        'language_display': LANGUAGE_DISPLAY,
        'presentation_templates': TEMPLATES,
        **pg,
    })
    return templates.TemplateResponse('admin/ai_presentations.html', context)


@router.get("/admin/ai/settings", response_class=HTMLResponse, dependencies=[Depends(require_admin)])
async def admin_ai_settings(request: Request, db: AsyncSession = Depends(get_db)):
    teachers = (await db.execute(
        select(Employee)
        .options(selectinload(Employee.school))
        .where(Employee.is_admin.is_(False))
        .order_by(Employee.first_name)
    )).scalars().all()

    counts = dict((await db.execute(
        select(AIPresentation.teacher_id, func.count(AIPresentation.id))
        .group_by(AIPresentation.teacher_id)
    )).all())

    q_counts = dict((await db.execute(
        select(AIQuestionSet.teacher_id, func.count(AIQuestionSet.id))
        .group_by(AIQuestionSet.teacher_id)
    )).all())

    context = await get_template_context(request, db)
    context.update({
        'teachers': teachers,
        'counts': counts,
        'q_counts': q_counts,
    })
    return templates.TemplateResponse('admin/ai_settings.html', context)


@router.post("/admin/ai/settings/{teacher_id}", dependencies=[Depends(require_admin)])
async def admin_ai_settings_save(
    teacher_id: int,
    request: Request,
    ai_enabled: Optional[str] = Form(None),
    ai_daily_limit: int = Form(3),
    ai_questions_daily_limit: int = Form(5),
    db: AsyncSession = Depends(get_db),
):
    teacher = (await db.execute(select(Employee).where(Employee.id == teacher_id))).scalar_one_or_none()
    if not teacher:
        raise HTTPException(404)

    teacher.ai_enabled = ai_enabled == 'on'
    teacher.ai_daily_limit = max(0, min(50, ai_daily_limit))
    teacher.ai_questions_daily_limit = max(0, min(100, ai_questions_daily_limit))
    await db.commit()

    flash(request, f"{teacher.first_name} {teacher.last_name} — {_t(request, 'ai_settings_saved_suffix')}", 'success')
    return RedirectResponse(url="/admin/ai/settings", status_code=303)


@router.get("/admin/ai/{pid}/download", dependencies=[Depends(require_admin)])
async def admin_ai_download(pid: int, db: AsyncSession = Depends(get_db)):
    presentation = (await db.execute(
        select(AIPresentation).where(AIPresentation.id == pid)
    )).scalar_one_or_none()
    if not presentation:
        raise HTTPException(404)

    try:
        data = json.loads(presentation.content)
        tpl_key = presentation.template if presentation.template in TEMPLATES else 'cosmos'
        pptx_bytes = await _build_pptx(data, presentation, TEMPLATES[tpl_key]['pptx'])
    except Exception as e:
        log.error("PPTX build failed for presentation %s: %s\n%s", pid, e, traceback.format_exc())
        raise HTTPException(500, detail=f"PPTX yaratishda xatolik: {e}")

    ascii_name = re.sub(r'[^A-Za-z0-9_\- ]+', '_', presentation.topic)[:60].strip('_ ') or 'presentation'
    ascii_filename = f"{ascii_name}_grade{presentation.grade}.pptx"
    utf8_filename = quote(f"{presentation.topic[:60]}_{presentation.grade}sinf.pptx")

    return StreamingResponse(
        BytesIO(pptx_bytes),
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={
            "Content-Disposition": f"attachment; filename=\"{ascii_filename}\"; filename*=UTF-8''{utf8_filename}",
            "Content-Length": str(len(pptx_bytes)),
        },
    )


# ============================================================================
# HELPERS
# ============================================================================

def _reset_if_new_day(teacher: Employee):
    today = date.today()
    if teacher.ai_last_reset != today:
        teacher.ai_used_today = 0
        teacher.ai_last_reset = today


async def _build_pptx(data: dict, presentation: AIPresentation, pptx_colors: dict) -> bytes:
    """Render presentation JSON into a PPTX file using template colors."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    C = pptx_colors  # shorthand

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]

    image_cache: dict = {}

    async def _fetch_one(client: httpx.AsyncClient, url: str):
        try:
            resp = await client.get(url, timeout=15.0)
            if resp.status_code == 200 and resp.headers.get('content-type', '').startswith('image/'):
                image_cache[url] = resp.content
        except Exception as e:
            log.warning("Image fetch failed %s: %s", url[:80], e)

    urls = [s['image_url'] for s in data.get('slides', []) if s.get('image_url')]
    if urls:
        async with httpx.AsyncClient(follow_redirects=True, timeout=15.0) as client:
            await asyncio.gather(*(_fetch_one(client, u) for u in urls), return_exceptions=True)

    def get_image(url: Optional[str]) -> Optional[BytesIO]:
        if not url or url not in image_cache:
            return None
        return BytesIO(image_cache[url])

    def add_text(slide, text: str, left, top, width, height, size=18, bold=False, color=None):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = text or ''
        run.font.size = Pt(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = RGBColor(*color)

    title_slide = prs.slides.add_slide(blank)
    add_text(title_slide, data.get('title', presentation.topic), Inches(0.5), Inches(2.5),
             Inches(12), Inches(1.5), size=44, bold=True, color=C['title'])
    add_text(title_slide, data.get('subtitle', ''), Inches(0.5), Inches(4.0),
             Inches(12), Inches(1), size=24, color=C['subtitle'])
    add_text(title_slide,
             f"{presentation.subject_name} · {presentation.grade}-sinf · {LANGUAGE_DISPLAY.get(presentation.language, presentation.language)}",
             Inches(0.5), Inches(6.5), Inches(12), Inches(0.5), size=14, color=C['muted'])

    for slide_data in data.get('slides', []):
        stype = slide_data.get('type', 'content')
        slide = prs.slides.add_slide(blank)

        if stype == 'title':
            add_text(slide, slide_data.get('title', ''), Inches(0.5), Inches(2.5), Inches(12),
                     Inches(1.5), size=40, bold=True, color=C['title'])
            add_text(slide, slide_data.get('subtitle', ''), Inches(0.5), Inches(4.0), Inches(12),
                     Inches(1), size=22, color=C['subtitle'])

        elif stype == 'content':
            add_text(slide, slide_data.get('title', ''), Inches(0.5), Inches(0.4), Inches(12),
                     Inches(0.9), size=30, bold=True, color=C['title'])
            points = slide_data.get('points', [])
            text = '\n'.join(f"•  {p}" for p in points)
            has_image = bool(slide_data.get('image_url'))
            text_width = Inches(6.5) if has_image else Inches(12)
            add_text(slide, text, Inches(0.5), Inches(1.5), text_width, Inches(5.5), size=18, color=C['body'])
            if has_image:
                img = get_image(slide_data.get('image_url'))
                if img:
                    try:
                        slide.shapes.add_picture(img, Inches(7.5), Inches(1.5), width=Inches(5.3))
                    except Exception as e:
                        log.warning("add_picture failed: %s", e)

        elif stype == 'image_focus':
            add_text(slide, slide_data.get('title', ''), Inches(0.5), Inches(0.4), Inches(12),
                     Inches(0.9), size=30, bold=True, color=C['title'])
            img = get_image(slide_data.get('image_url'))
            if img:
                try:
                    slide.shapes.add_picture(img, Inches(2.5), Inches(1.5), width=Inches(8.3))
                except Exception as e:
                    log.warning("add_picture failed: %s", e)
            add_text(slide, slide_data.get('caption', ''), Inches(0.5), Inches(6.5), Inches(12),
                     Inches(0.7), size=16, color=C['subtitle'])

        elif stype == 'quiz':
            add_text(slide, slide_data.get('question', ''), Inches(0.5), Inches(0.5),
                     Inches(12), Inches(1.5), size=26, bold=True, color=C['title'])
            options = slide_data.get('options', [])
            correct = slide_data.get('correct', 0)
            for i, opt in enumerate(options):
                marker = "✓" if i == correct else "•"
                color = C['correct'] if i == correct else C['body']
                add_text(slide, f"  {marker}  {opt}", Inches(1), Inches(2.3 + i * 0.7),
                         Inches(11), Inches(0.6), size=20, color=color)
            add_text(slide, slide_data.get('explanation', ''), Inches(0.5), Inches(5.8),
                     Inches(12), Inches(1.2), size=14, color=C['muted'])

        elif stype == 'true_false':
            add_text(slide, slide_data.get('statement', ''), Inches(0.5), Inches(2),
                     Inches(12), Inches(2), size=28, bold=True, color=C['title'])
            answer = "✓ TO'G'RI" if slide_data.get('is_true') else "✗ NOTO'G'RI"
            color = C['correct'] if slide_data.get('is_true') else C['wrong']
            add_text(slide, answer, Inches(0.5), Inches(4.2), Inches(12), Inches(1),
                     size=32, bold=True, color=color)
            add_text(slide, slide_data.get('explanation', ''), Inches(0.5), Inches(5.5),
                     Inches(12), Inches(1.5), size=14, color=C['muted'])

        elif stype == 'fill_blank':
            add_text(slide, slide_data.get('title', ''), Inches(0.5), Inches(0.4), Inches(12),
                     Inches(0.9), size=28, bold=True, color=C['title'])
            sentence = slide_data.get('sentence', '').replace('___', f"[{slide_data.get('answer', '____')}]")
            add_text(slide, sentence, Inches(0.5), Inches(2.5), Inches(12), Inches(2), size=24, color=C['body'])
            if slide_data.get('hint'):
                add_text(slide, slide_data['hint'], Inches(0.5), Inches(5.5),
                         Inches(12), Inches(1), size=14, color=C['muted'])

        elif stype == 'match':
            add_text(slide, slide_data.get('title', ''), Inches(0.5), Inches(0.4), Inches(12),
                     Inches(0.9), size=28, bold=True, color=C['title'])
            for i, pair in enumerate(slide_data.get('pairs', [])):
                if len(pair) >= 2:
                    add_text(slide, str(pair[0]), Inches(1), Inches(1.7 + i * 0.7), Inches(5),
                             Inches(0.6), size=18, bold=True, color=C['title'])
                    add_text(slide, "→  " + str(pair[1]), Inches(7), Inches(1.7 + i * 0.7),
                             Inches(5), Inches(0.6), size=18, color=C['body'])

        elif stype == 'sequence':
            add_text(slide, slide_data.get('title', ''), Inches(0.5), Inches(0.4), Inches(12),
                     Inches(0.9), size=28, bold=True, color=C['title'])
            for i, item in enumerate(slide_data.get('items', [])):
                add_text(slide, f"{i + 1}.  {item}", Inches(1), Inches(1.7 + i * 0.7),
                         Inches(11), Inches(0.6), size=20, color=C['body'])
            if slide_data.get('explanation'):
                add_text(slide, slide_data['explanation'], Inches(0.5), Inches(5.8),
                         Inches(12), Inches(1.2), size=14, color=C['muted'])

        elif stype == 'open_question':
            add_text(slide, slide_data.get('question', ''), Inches(0.5), Inches(1.5),
                     Inches(12), Inches(2.5), size=30, bold=True, color=C['title'])
            hints = slide_data.get('hints', [])
            if hints:
                hint_text = '\n'.join(f"•  {h}" for h in hints)
                add_text(slide, hint_text, Inches(0.5), Inches(4.5), Inches(12), Inches(2.5),
                         size=16, color=C['subtitle'])

        elif stype == 'summary':
            add_text(slide, slide_data.get('title', 'Xulosa'), Inches(0.5), Inches(0.4),
                     Inches(12), Inches(0.9), size=30, bold=True, color=C['title'])
            text = '\n'.join(f"•  {p}" for p in slide_data.get('points', []))
            add_text(slide, text, Inches(0.5), Inches(1.7), Inches(12), Inches(5), size=20, color=C['body'])

        else:
            add_text(slide, slide_data.get('title', ''), Inches(0.5), Inches(0.4), Inches(12),
                     Inches(0.9), size=28, bold=True, color=C['title'])
            add_text(slide, json.dumps(slide_data, ensure_ascii=False), Inches(0.5), Inches(1.5),
                     Inches(12), Inches(5), size=14, color=C['body'])

    out = BytesIO()
    prs.save(out)
    return out.getvalue()
